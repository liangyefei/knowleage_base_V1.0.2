#!/usr/bin/env python3
# @Time : 2024/4/9 10:30:37
# @Author : Hetao
# @File : [Code] --> myexcelloader.py
# @Software: PyCharm
# @Function: excel解析增强 - 使用大模型解析Excel结构

import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
import requests
from jinja2 import Template
from langchain.document_loaders import UnstructuredFileLoader
from openpyxl.utils import get_column_letter

from configs.logger_manager import logger
from configs.model_loader import get_config

# ✅ 导入 prompt 加载器
from configs.prompt_loader import get_prompt

from document_loaders.base_loader import _save_chunks_to_json

"""
    Excel加载器，用于加载Excel文件。
    TODO CK  注意参数从configs导入。
"""

ONLINE_LLM_CONFIG = get_config("ONLINE_LLM_MODEL", {})
MODEL_CONFIG = ONLINE_LLM_CONFIG.get("deepseek-v4-pro", {})

LLM_CONFIG = {
    "api_key": MODEL_CONFIG.get("api_key"),
    "base_url": MODEL_CONFIG.get("api_base_url"),
    "model": MODEL_CONFIG.get("model_id"),
    "max_tokens": 8192,
    "temperature": 0.1
}

if MODEL_CONFIG:
    logger.info(f"已加载模型配置: {MODEL_CONFIG.get('description', '')}")
else:
    logger.warning("未找到 ONLINE_LLM_MODEL 配置，请检查 model_config.yaml")

_CONTENT_KEY = 'content'


def call_llm(prompt):
    """
    调用百炼平台的大模型 API（兼容 OpenAI 接口）
    支持超时重试，最多重试2次
    """
    _DEFAULT_MAX_RETRIES = 2
    _INITIAL_RETRY_COUNT = 0
    _DEFAULT_BASE_DELAY = 5
    _CHOICES_KEY = 'choices'
    # 检查配置是否完整
    if not LLM_CONFIG.get("api_key") or not LLM_CONFIG.get("base_url"):
        logger.error("LLM 配置不完整，请检查 model_config.yaml 中的 ONLINE_LLM_MODEL 配置")
        return None

    url = f"{LLM_CONFIG['base_url']}/chat/completions"
    headers = {
        "Authorization": f"Bearer {LLM_CONFIG['api_key']}",
        "Content-Type": "application/json"
    }
    payload = {
        "model": LLM_CONFIG["model"],
        "messages": [
            {"role": "system", _CONTENT_KEY: get_prompt("excel_detection", "system_prompt")},
            {"role": "user", _CONTENT_KEY: prompt}
        ],
        "max_tokens": LLM_CONFIG["max_tokens"],
        "temperature": LLM_CONFIG["temperature"]
    }

    max_retries = _DEFAULT_MAX_RETRIES
    retry_count = _INITIAL_RETRY_COUNT
    base_delay = _DEFAULT_BASE_DELAY

    while retry_count <= max_retries:
        try:
            if retry_count > 0:
                logger.info(f"  重试第 {retry_count} 次...")

            response = requests.post(url, headers=headers, json=payload, timeout=(30, 120))

            if response.status_code == 200:
                result = response.json()

                if _CHOICES_KEY not in result or len(result[_CHOICES_KEY]) == 0:
                    logger.error(f"  API返回结构异常: {result}")
                    return None

                content = result[_CHOICES_KEY][0].get("message", {}).get(_CONTENT_KEY, "")

                if not content or content.strip() == "":
                    logger.error("  API返回内容为空")
                    return None

                logger.info(f"  大模型返回: {content[:200]}...")
                return content

            elif response.status_code == 429:
                retry_count += 1
                if retry_count <= max_retries:
                    wait_time = base_delay * (retry_count + 1)
                    logger.error(f"  API限流，等待 {wait_time} 秒后重试... (已重试 {retry_count}/{max_retries})")
                    time.sleep(wait_time)
                else:
                    logger.error(f"  API限流，已达到最大重试次数 ({max_retries})，放弃")
                    return None
            else:
                logger.error(f"  API调用失败，状态码: {response.status_code}")
                logger.error(f"  错误信息: {response.text}")
                return None

        except requests.exceptions.Timeout:
            retry_count += 1
            if retry_count <= max_retries:
                wait_time = base_delay * (retry_count + 1)
                logger.error(f"  API调用超时，等待 {wait_time} 秒后重试... (已重试 {retry_count}/{max_retries})")
                time.sleep(wait_time)
            else:
                logger.error(f"  API调用超时，已达到最大重试次数 ({max_retries})，放弃")
                return None

        except requests.exceptions.ConnectionError as e:
            retry_count += 1
            if retry_count <= max_retries:
                wait_time = base_delay * (retry_count + 1)
                logger.error(f"  网络连接错误，等待 {wait_time} 秒后重试... (已重试 {retry_count}/{max_retries})")
                time.sleep(wait_time)
            else:
                logger.error(f"  网络连接错误，已达到最大重试次数 ({max_retries})，放弃: {e}")
                return None

        except json.JSONDecodeError as e:
            logger.error(f"  JSON解析失败: {e}")
            logger.error(f"  原始响应: {response.text[:300] if 'response' in locals() else '无法获取'}")
            return None

        except Exception as e:
            retry_count += 1
            if retry_count <= max_retries:
                wait_time = base_delay * (retry_count + 1)
                logger.error(f"  API调用异常，等待 {wait_time} 秒后重试... (已重试 {retry_count}/{max_retries}): {e}")
                time.sleep(wait_time)
            else:
                logger.error(f"  API调用异常，已达到最大重试次数 ({max_retries})，放弃: {e}")
                return None

    return None


# ============================================================
# 第一步：用大模型检测说明文本的位置和表名（完整保留你的代码）
# ============================================================

def detect_description_and_table_name_with_llm(ws):
    """
    用大模型检测 sheet 中哪些行是说明文本（非表格数据），同时检测表名
    返回：{
        'data_start_row': int,
        'data_end_row': int,
        'table_name': str,
        'table_name_row': int or None
    }
    """
    _DETECT_DESCRIPT_LABEL = '```'
    # 读取前8行内容
    top_text = []
    for row_idx in range(1, min(9, ws.max_row + 1)):
        row_values = []
        for col_idx in range(1, min(10, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                row_values.append(str(cell.value).strip())
        top_text.append(f"第{row_idx}行: {' | '.join(row_values)}")

    # 读取倒数5行内容
    bottom_text = []
    start_row = max(1, ws.max_row - 4)
    for row_idx in range(start_row, ws.max_row + 1):
        row_values = []
        for col_idx in range(1, min(10, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value:
                row_values.append(str(cell.value).strip())
        bottom_text.append(f"第{row_idx}行: {' | '.join(row_values)}")

    # ✅ 从配置文件加载提示词模板
    user_template_str = get_prompt("excel_detection", "user_prompt")
    template = Template(user_template_str)
    prompt = template.render(
        top_text='\n'.join(top_text),
        bottom_text='\n'.join(bottom_text)
    )

    response = call_llm(prompt)
    if response:
        try:
            if "```json" in response:
                response = response.split("```json")[1].split(_DETECT_DESCRIPT_LABEL)[0].strip()
            elif _DETECT_DESCRIPT_LABEL in response:
                response = response.split(_DETECT_DESCRIPT_LABEL)[1].split(_DETECT_DESCRIPT_LABEL)[0].strip()

            result = json.loads(response)
            return result
        except json.JSONDecodeError as e:
            logger.error(f"  大模型返回JSON解析失败: {e}")
            logger.error(f"  原始返回: {response[:300]}...")
            return None
    else:
        return None


# ============================================================
# 第二步：检测表头行数（完整保留你的代码）
# ============================================================

def detect_header_rows_with_llm(ws, max_check_rows=8):
    """
    用大模型检测 sheet 中哪些行是表头行
    返回：表头行号列表，如 [2, 3] 表示第2行和第3行是表头
    """
    _DETECT_HEADER_R_LABEL = '```'
    # 读取前 max_check_rows 行内容
    top_text = []
    for row_idx in range(1, min(max_check_rows + 1, ws.max_row + 1)):
        row_values = []
        for col_idx in range(1, min(15, ws.max_column + 1)):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value is not None and str(cell.value).strip():
                row_values.append(str(cell.value).strip())
        top_text.append(f"第{row_idx}行: {' | '.join(row_values)}")

    # ✅ 从配置文件加载提示词模板
    user_template_str = get_prompt("excel_header_detection", "user_prompt")
    template = Template(user_template_str)
    prompt = template.render(
        max_check_rows=max_check_rows,
        top_text='\n'.join(top_text)
    )

    response = call_llm(prompt)
    if response:
        try:
            if "```json" in response:
                response = response.split("```json")[1].split(_DETECT_HEADER_R_LABEL)[0].strip()
            elif _DETECT_HEADER_R_LABEL in response:
                response = response.split(_DETECT_HEADER_R_LABEL)[1].split(_DETECT_HEADER_R_LABEL)[0].strip()

            result = json.loads(response)
            header_rows = result.get('header_rows', [])
            if header_rows and isinstance(header_rows, list):
                return header_rows
            else:
                return []
        except json.JSONDecodeError as e:
            logger.error(f"  大模型返回JSON解析失败: {e}")
            logger.error(f"  原始返回: {response[:300]}...")
            return []
    else:
        return []


# ============================================================
# 第三步：拆分合并单元格 + 填充（完整保留你的代码）
# ============================================================

def unmerge_and_fill_sheet(ws):
    """
    拆分指定sheet的所有合并单元格，并用原值填充
    """
    # 第一步：收集所有合并区域的信息
    merge_info = []
    for merged_range in ws.merged_cells.ranges:
        min_row = merged_range.min_row
        min_col = merged_range.min_col
        max_row = merged_range.max_row
        max_col = merged_range.max_col
        top_left_value = ws.cell(row=min_row, column=min_col).value

        merge_info.append({
            'min_row': min_row,
            'min_col': min_col,
            'max_row': max_row,
            'max_col': max_col,
            'value': top_left_value
        })

    if not merge_info:
        return 0

    # 第二步：拆分所有合并单元格
    merge_ranges = list(ws.merged_cells.ranges)
    for merged_range in merge_ranges:
        try:
            ws.unmerge_cells(
                start_row=merged_range.min_row,
                start_column=merged_range.min_col,
                end_row=merged_range.max_row,
                end_column=merged_range.max_col
            )
        except Exception:
            try:
                range_str = f"{get_column_letter(merged_range.min_col)}{merged_range.min_row}:{get_column_letter(merged_range.max_col)}{merged_range.max_row}"  # pylint: disable=line-too-long
                ws.unmerge_cells(range_str)
            except Exception as e:
                logger.error(f"    拆分失败: {range_str}, 错误: {e}")

    # 第三步：填充空值
    fill_count = 0
    for info in merge_info:
        for row in range(info['min_row'], info['max_row'] + 1):
            for col in range(info['min_col'], info['max_col'] + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value is None:
                    cell.value = info['value']
                    fill_count += 1

    return fill_count


# ============================================================
# 第四步：合并多行表头 + 删除"列X"列（完整保留你的代码）
# ============================================================

def merge_multi_row_header(ws, header_rows):
    """
    合并多行表头为单行，并删除所有表头为"列X"的列
    header_rows: 表头行号列表，如 [2, 3]
    """
    if not header_rows or len(header_rows) == 1:
        return 0

    header_rows = sorted(header_rows)
    logger.info(f"  合并表头行: {header_rows}")

    # ---- 第1步：获取所有表头行的内容 ----
    header_data = []
    for row_idx in header_rows:
        row_values = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = cell.value if cell.value is not None else ""
            row_values.append(str(value).strip())
        header_data.append(row_values)

    # ---- 第2步：按列合并表头 ----
    merged_header = []
    for col_idx in range(ws.max_column):
        col_values = []
        for row_idx in range(len(header_rows)):
            val = header_data[row_idx][col_idx]
            if val and val != "":
                col_values.append(val)

        if not col_values:
            merged_header.append("")
        elif len(col_values) == 1:
            merged_header.append(col_values[0])
        else:
            unique_values = []
            for v in col_values:
                if v not in unique_values:
                    unique_values.append(v)
            if len(unique_values) == 1:
                merged_header.append(unique_values[0])
            else:
                merged_header.append("_".join(unique_values))

    logger.info(f"  合并后的表头: {merged_header[:10]}..." if len(merged_header) > 10 else f"  合并后的表头: {merged_header}")

    # ---- 第3步：找出所有以"列"开头的表头列 ----
    cols_to_delete = []
    for col_idx, header_value in enumerate(merged_header, start=1):
        if header_value and header_value.startswith("列"):
            cols_to_delete.append(col_idx)
            logger.info(f"    发现无用列: 第{col_idx}列, 表头: '{header_value}'")

    # ---- 第4步：获取所有数据行 ----
    data_start_row = max(header_rows) + 1
    data_rows = []
    for row_idx in range(data_start_row, ws.max_row + 1):
        row_values = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            row_values.append(cell.value)
        data_rows.append(row_values)

    # ---- 第4.5步：过滤空行（只过滤空行，不过滤重复行） ----
    filtered_data_rows = []
    for row_data in data_rows:
        # 检查是否为空行：所有值都为空
        non_empty = [v for v in row_data if v is not None and str(v).strip() != ""]
        if non_empty:
            filtered_data_rows.append(row_data)

    if len(filtered_data_rows) < len(data_rows):
        logger.info(f"  过滤掉 {len(data_rows) - len(filtered_data_rows)} 个空行")
    data_rows = filtered_data_rows

    # ---- 第5步：删除"列X"对应的列（从右往左删） ----
    if cols_to_delete:
        logger.info(f"  准备删除 {len(cols_to_delete)} 个无用列: {cols_to_delete}")

        for col_idx in sorted(cols_to_delete, reverse=True):
            merged_header.pop(col_idx - 1)

        for row in data_rows:
            for col_idx in sorted(cols_to_delete, reverse=True):
                if col_idx - 1 < len(row):
                    row.pop(col_idx - 1)

        for col_idx in sorted(cols_to_delete, reverse=True):
            ws.delete_cols(col_idx)

        logger.info(f"  已删除 {len(cols_to_delete)} 个无用列，剩余 {len(merged_header)} 列")

    # ---- 第6步：删除所有表头行以及表头之前的行 ----
    rows_to_delete = []
    min_header_row = min(header_rows)
    for row_idx in range(1, min_header_row):
        rows_to_delete.append(row_idx)
    for row_idx in header_rows:
        rows_to_delete.append(row_idx)

    for row_idx in sorted(rows_to_delete, reverse=True):
        ws.delete_rows(row_idx)

    logger.info(f"  已删除 {len(rows_to_delete)} 行（含表头及之前的行）")

    # ---- 第7步：清空所有剩余行 ----
    while ws.max_row > 1:
        ws.delete_rows(ws.max_row)

    # ---- 第8步：在第1行插入合并后的表头 ----
    ws.insert_rows(1)
    for col_idx, header_value in enumerate(merged_header, start=1):
        ws.cell(row=1, column=col_idx).value = header_value

    # ---- 第9步：数据行从第2行开始写入 ----
    for data_row_idx, row_data in enumerate(data_rows, start=2):
        non_empty = [v for v in row_data if v is not None and str(v).strip() != ""]
        if not non_empty:
            continue
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=data_row_idx, column=col_idx).value = value

    logger.info(f"  已合并 {len(header_rows)} 行表头为单行，共 {len(data_rows)} 行数据，剩余 {len(merged_header)} 列")
    return len(header_rows)


# ============================================================
# 第五步：删除空列（完整保留你的代码）
# ============================================================

def delete_empty_columns(ws):
    """
    删除 sheet 中所有完全为空的列
    使用 iter_cols 批量扫描 + 合并连续区间批量删除
    返回删除的列数
    """
    if ws.max_column == 0:
        return 0

    # ---- 用 iter_cols 批量扫描空列 ----
    empty_cols = []
    for col_idx, col_cells in enumerate(
        ws.iter_cols(min_col=1, max_col=ws.max_column, values_only=True), start=1
    ):
        is_empty = True
        for v in col_cells:
            if v is not None and str(v).strip() != '':
                is_empty = False
                break
        if is_empty:
            empty_cols.append(col_idx)

    if not empty_cols:
        return 0

    # ---- 合并连续空列区间 ----
    ranges = []  # [(start, count), ...]
    start = empty_cols[0]
    count = 1
    for i in range(1, len(empty_cols)):
        if empty_cols[i] == empty_cols[i - 1] + 1:
            count += 1
        else:
            ranges.append((start, count))
            start = empty_cols[i]
            count = 1
    ranges.append((start, count))

    logger.info(f"    发现 {len(empty_cols)} 个空列，合并为 {len(ranges)} 个连续区间: {ranges}")

    # ---- 从后往前批量删除 ----
    for start, count in sorted(ranges, reverse=True):
        ws.delete_cols(start, count)

    logger.info(f"    已删除 {len(empty_cols)} 个空列，剩余 {ws.max_column} 列")

    return len(empty_cols)


# ============================================================
# 第六步：Excel内存分块器（不保存文件）
# ============================================================

class ExcelChunkerInMemory:
    """
    内存版Excel分块器 - 不保存任何文件
    """

    def __init__(self, wb, metadata: dict, file_name: str):
        self.wb = wb
        self.metadata = metadata
        self.file_name = f"{file_name}.xlsx"

    @staticmethod
    def get_headers(ws) -> list[str]:
        headers = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            value = cell.value if cell.value is not None else ""
            headers.append(str(value).strip())
        return headers

    @staticmethod
    def _convert_value(value):
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        if hasattr(value, 'isoformat'):
            try:
                return value.isoformat()
            except (TypeError, ValueError):
                return str(value)
        return value

    def row_to_dict(self, ws, row_idx: int, headers: list[str]) -> dict[str, Any]:
        row_dict = {}
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            value = self._convert_value(cell.value)
            if not header:
                header = f"列{col_idx}"
            row_dict[header] = value
        return row_dict

    def rows_to_markdown_table(self, ws, start_row: int, end_row: int, headers: list[str]) -> str:
        """
        将指定行范围的数据转为Markdown表格字符串
        每个 chunk 都自带表头行
        """
        _ROWS_TO_MARKDOW_LABEL = '| '
        _MD_COL_DELIMITER = ' | '
        _MD_ROW_END = ' |'
        _MD_SEPARATOR_CELL = '---'
        lines = []

        # 表头行
        header_cells = [str(h).replace('|', '\\|').replace('\n', ' ') if h else '' for h in headers]
        lines.append(_ROWS_TO_MARKDOW_LABEL + _MD_COL_DELIMITER.join(header_cells) + _MD_ROW_END)

        # 分隔行
        lines.append(_ROWS_TO_MARKDOW_LABEL + _MD_COL_DELIMITER.join([_MD_SEPARATOR_CELL] * len(headers)) + _MD_ROW_END)

        # 数据行
        for row_idx in range(start_row, end_row + 1):
            row_cells = []
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = self._convert_value(cell.value)
                if isinstance(value, str):
                    value = value.replace('|', '\\|').replace('\n', ' ').strip()
                row_cells.append(str(value))
            lines.append(_ROWS_TO_MARKDOW_LABEL + _MD_COL_DELIMITER.join(row_cells) + _MD_ROW_END)

        return '\n'.join(lines)

    def process_sheet_blocks(self, ws, sheet_name: str, sheet_meta: dict) -> list[dict]:
        _BLOCK_ID_KEY = 'block_id'
        _CONTENT_KEY = 'content'
        _METADATA_KEY = 'metadata'
        _FILE_NAME_KEY = 'file_name'
        _SHEET_NAME_KEY = 'sheet_name'
        _HEADERS_KEY = 'headers'
        _CONTENT_TYPE_KEY = 'content_type'
        blocks = []

        SKIP_SHEETS = ['版权声明', '0.说明', '版权申明']
        if sheet_name in SKIP_SHEETS:
            content_lines = []
            for row_idx in range(1, ws.max_row + 1):
                row_values = []
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if cell.value is not None:
                        row_values.append(str(self._convert_value(cell.value)).strip())
                if row_values:
                    content_lines.append(' '.join(row_values))

            if content_lines:
                blocks.append({
                    _BLOCK_ID_KEY: '',
                    _CONTENT_KEY: '\n'.join(content_lines),
                    _METADATA_KEY: {
                        _FILE_NAME_KEY: self.file_name,
                        _SHEET_NAME_KEY: sheet_name,
                        _HEADERS_KEY: [],
                        'skipped_sheet': True
                    }
                })
            return blocks

        header_rows_detected = sheet_meta.get('header_rows_detected', [])
        description_before = sheet_meta.get('description_before', '')
        description_after = sheet_meta.get('description_after', '')
        table_name = sheet_meta.get('table_name', '')
        is_pure_description = sheet_meta.get('is_pure_description', False)

        if is_pure_description:
            content = sheet_meta.get(_CONTENT_KEY, '')
            if content:
                blocks.append({
                    _BLOCK_ID_KEY: '',
                    _CONTENT_KEY: content,
                    _METADATA_KEY: {
                        _FILE_NAME_KEY: self.file_name,
                        _SHEET_NAME_KEY: sheet_name,
                        _HEADERS_KEY: [],
                        _CONTENT_TYPE_KEY: 'pure_description'
                    }
                })
            return blocks

        if header_rows_detected:
            headers = self.get_headers(ws)
            valid_headers = [h for h in headers if h and h != "" and not h.startswith("列")]

            if valid_headers:
                # ============ 滑动窗口分块（不重叠）============
                _CHUNK_SIZE_PER_BLOCK = 15
                data_start = 2
                data_rows = []
                for row_idx in range(data_start, ws.max_row + 1):
                    row_dict = self.row_to_dict(ws, row_idx, headers)
                    non_empty_values = [v for v in row_dict.values() if v and v != ""]
                    if not non_empty_values:
                        continue
                    data_rows.append((row_idx, row_dict))

                if not data_rows:
                    return blocks

                total_rows = len(data_rows)
                chunk_count = (total_rows + _CHUNK_SIZE_PER_BLOCK - 1) // _CHUNK_SIZE_PER_BLOCK
                logger.info(f"  表格共 {total_rows} 行数据，分 {chunk_count} 块（每块{_CHUNK_SIZE_PER_BLOCK}行）")

                for i in range(chunk_count):
                    start = i * _CHUNK_SIZE_PER_BLOCK
                    end = min(start + _CHUNK_SIZE_PER_BLOCK, total_rows)
                    actual_start_row = data_rows[start][0]
                    actual_end_row = data_rows[end - 1][0]

                    # 构建Markdown表格
                    md_table = self.rows_to_markdown_table(
                        ws, actual_start_row, actual_end_row, headers
                    )

                    # 拼接表名和说明文本
                    content_parts = []
                    if table_name:
                        content_parts.append(f"表名：{table_name}")
                    description_text = description_before or description_after
                    if description_text:
                        content_parts.append(f"说明：{description_text}")
                    content_parts.append(md_table)

                    blocks.append({
                        _BLOCK_ID_KEY: '',
                        _CONTENT_KEY: '\n\n'.join(content_parts),
                        _METADATA_KEY: {
                            _FILE_NAME_KEY: self.file_name,
                            _SHEET_NAME_KEY: sheet_name,
                            _HEADERS_KEY: headers,
                            'row_range': f"{actual_start_row}-{actual_end_row}"
                        }
                    })
            else:
                # 无有效表头，按原逻辑逐行处理
                headers = self.get_headers(ws)
                for row_idx in range(1, ws.max_row + 1):
                    row_dict = self.row_to_dict(ws, row_idx, headers)
                    non_empty_values = [v for v in row_dict.values() if v and v != ""]
                    if not non_empty_values:
                        continue
                    blocks.append({
                        _BLOCK_ID_KEY: '',
                        _CONTENT_KEY: json.dumps(row_dict, ensure_ascii=False),
                        _METADATA_KEY: {
                            _FILE_NAME_KEY: self.file_name,
                            _SHEET_NAME_KEY: sheet_name,
                            _HEADERS_KEY: headers,
                            'row_number': row_idx
                        }
                    })
        else:
            if description_before:
                blocks.append({
                    _BLOCK_ID_KEY: '',
                    _CONTENT_KEY: description_before,
                    _METADATA_KEY: {
                        _FILE_NAME_KEY: self.file_name,
                        _SHEET_NAME_KEY: sheet_name,
                        _HEADERS_KEY: [],
                        _CONTENT_TYPE_KEY: 'description_before'
                    }
                })

            if description_after:
                blocks.append({
                    _BLOCK_ID_KEY: '',
                    _CONTENT_KEY: description_after,
                    _METADATA_KEY: {
                        _FILE_NAME_KEY: self.file_name,
                        _SHEET_NAME_KEY: sheet_name,
                        _HEADERS_KEY: [],
                        _CONTENT_TYPE_KEY: 'description_after'
                    }
                })

            if table_name and not description_before and not description_after:
                blocks.append({
                    _BLOCK_ID_KEY: '',
                    _CONTENT_KEY: f"表名: {table_name}",
                    _METADATA_KEY: {
                        _FILE_NAME_KEY: self.file_name,
                        _SHEET_NAME_KEY: sheet_name,
                        _HEADERS_KEY: [],
                        _CONTENT_TYPE_KEY: 'table_name'
                    }
                })

        return blocks

    def chunk_excel(self) -> list[dict[str, Any]]:
        all_blocks = []

        for sheet_name in self.wb.sheetnames:
            ws = self.wb[sheet_name]
            sheet_meta = self.metadata.get(sheet_name, {})

            logger.info(f"\n处理 Sheet: {sheet_name}")
            logger.info(f"  行数: {ws.max_row}, 列数: {ws.max_column}")

            blocks = self.process_sheet_blocks(ws, sheet_name, sheet_meta)
            logger.info(f"  生成 {len(blocks)} 个块")
            all_blocks.extend(blocks)

        for idx, block in enumerate(all_blocks, start=1):
            block['block_id'] = f"{idx:08d}"

        return all_blocks


# ============================================================
# ExcelChineseLoader - 主加载器（接口保持不变，支持并发）
# ============================================================

class ExcelChineseLoader(UnstructuredFileLoader):
    """
    Excel加载器 - 使用大模型解析Excel结构
    支持并发处理多个Excel文件
    接口与原来完全兼容
    """

    def __init__(
            self,
            file_path: str,
            max_workers: int = 1,
            **kwargs
    ):
        super().__init__(file_path, **kwargs)
        self.max_workers = max_workers

    from langchain_core.documents import Document

    @staticmethod
    def _convert_pandas_to_openpyxl(df_dict: dict[str, pd.DataFrame]) -> openpyxl.Workbook:
        """
        将 pandas DataFrame 字典转换为 openpyxl Workbook 对象
        保持与原有代码的兼容性
        """
        wb = openpyxl.Workbook()

        # 删除默认创建的 Sheet
        default_sheet = wb.active
        if default_sheet:
            wb.remove(default_sheet)

        for sheet_name, df in df_dict.items():
            ws = wb.create_sheet(title=sheet_name)

            # 写入表头
            for col_idx, col_name in enumerate(df.columns, start=1):
                ws.cell(row=1, column=col_idx).value = col_name

            # 写入数据
            for row_idx, row_data in enumerate(df.values, start=2):
                for col_idx, value in enumerate(row_data, start=1):
                    # 处理 NaN 和 NaT
                    if pd.isna(value):
                        continue
                    # 处理日期时间
                    if isinstance(value, (pd.Timestamp, pd.DatetimeIndex)):
                        value = value.to_pydatetime()
                    ws.cell(row=row_idx, column=col_idx).value = value

        return wb

    @classmethod
    def process_batch(cls, file_paths: list[str], max_workers: int = 3, **kwargs) -> dict[str, list]:
        """
        批量并发处理多个Excel文件（新增方法）

        Args:
            file_paths: Excel文件路径列表
            max_workers: 最大并发数
            **kwargs: 传递给 ExcelChineseLoader 的参数

        Returns:
            Dict[str, List]: {file_path: elements列表}
        """
        results = {}
        total = len(file_paths)
        completed = 0

        logger.info(f"\n{'=' * 60}")
        logger.info(f"开始批量处理 {total} 个Excel文件，并发数: {max_workers}")
        logger.info(f"{'=' * 60}")

        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            future_to_path = {}
            for file_path in file_paths:
                loader = cls(file_path=file_path, max_workers=1, **kwargs)
                future = executor.submit(loader._process_single_excel, file_path)
                future_to_path[future] = file_path

            for future in as_completed(future_to_path):
                file_path = future_to_path[future]
                completed += 1
                try:
                    path, elements = future.result()
                    results[path] = elements
                    logger.info(f"\n📊 进度: {completed}/{total} - {path} 完成")
                except Exception as e:
                    logger.error(f"\n❌ {file_path} 处理失败: {e}")
                    results[file_path] = []

        logger.info(f"\n{'=' * 60}")
        logger.info(f"批量处理完成！成功: {len([r for r in results.values() if r])} 个")
        logger.info(f"{'=' * 60}")

        return results

    def load(self) -> list[Document]:
        """重写 load 方法，直接返回 _get_elements 的结果，避免父类合并"""
        return self._get_elements()

    def _process_single_excel(self, file_path: str) -> tuple[str, list]:
        """处理单个Excel文件 - 不保存任何中间文件"""
        _TABLE_NAME_KEY = 'table_name'
        _TABLE_NAME_ROW_KEY = 'table_name_row'
        try:
            logger.info(f"\n{'=' * 60}")
            logger.info(f"开始解析Excel: {file_path}")
            logger.info(f"{'=' * 60}")

            file_ext = Path(file_path).suffix.lower()
            logger.info(f"文件格式: {file_ext}")

            # ---- 根据扩展名选择加载方式 ----
            if file_ext == '.xls':
                # 使用 pandas + xlrd 读取 .xls 格式
                logger.info("📄 使用 xlrd 引擎加载 .xls 文件")
                df_dict = pd.read_excel(file_path, sheet_name=None, engine='xlrd')
                # 转换为 openpyxl Workbook 对象
                wb = self._convert_pandas_to_openpyxl(df_dict)

            elif file_ext == '.xlsx':
                # 使用 openpyxl 读取 .xlsx 格式（原有逻辑）
                logger.info("📄 使用 openpyxl 加载 .xlsx 文件")
                wb = openpyxl.load_workbook(file_path, data_only=True)

            else:
                raise ValueError(f"不支持的文件格式: {file_ext}，仅支持 .xlsx 和 .xls")

            logger.info(f"共有 {len(wb.sheetnames)} 个Sheet: {wb.sheetnames}")

            metadata = {}
            SKIP_SHEETS = ['版权声明', '0.说明', '版权申明']

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                logger.info(f"\n{'=' * 50}")
                logger.info(f"正在处理 Sheet: {sheet_name}")
                logger.info(f"  原始行数: {ws.max_row}, 列数: {ws.max_column}")

                # ---- 判断是否为跳过的sheet ----
                if sheet_name in SKIP_SHEETS:
                    logger.info(f"  该Sheet为'{sheet_name}'，跳过处理，仅提取元数据")
                    content_lines = []
                    for row_idx in range(1, min(ws.max_row + 1, 50)):
                        row_values = []
                        for col_idx in range(1, ws.max_column + 1):
                            cell = ws.cell(row=row_idx, column=col_idx)
                            if cell.value is not None:
                                row_values.append(str(cell.value).strip())
                        if row_values:
                            content_lines.append(' '.join(row_values))

                    metadata[sheet_name] = {
                        'skipped': True,
                        'reason': f'Sheet名称为"{sheet_name}"，跳过处理',
                        _CONTENT_KEY: '\n'.join(content_lines) if content_lines else ''
                    }
                    continue

                # ---- 正常处理流程 ----
                result = detect_description_and_table_name_with_llm(ws)

                if result is None:
                    logger.info("  无法检测表格结构，保留全部内容")
                    data_start_row = 1
                    data_end_row = ws.max_row
                    table_name = ""
                    table_name_row = None
                else:
                    data_start_row = result.get('data_start_row')
                    data_end_row = result.get('data_end_row')
                    table_name = result.get(_TABLE_NAME_KEY, '')
                    table_name_row = result.get(_TABLE_NAME_ROW_KEY)

                # 检查是否为纯说明文本
                if data_start_row is None and data_end_row is None:
                    logger.info("  该Sheet为纯说明文本（无表格数据），跳过处理")
                    content_lines = []
                    for row_idx in range(1, ws.max_row + 1):
                        row_values = []
                        for col_idx in range(1, ws.max_column + 1):
                            cell = ws.cell(row=row_idx, column=col_idx)
                            if cell.value is not None:
                                row_values.append(str(cell.value).strip())
                        if row_values:
                            content_lines.append(' '.join(row_values))

                    metadata[sheet_name] = {
                        'is_pure_description': True,
                        _CONTENT_KEY: '\n'.join(content_lines) if content_lines else '',
                        _TABLE_NAME_KEY: table_name,
                        _TABLE_NAME_ROW_KEY: table_name_row
                    }
                    continue

                # 兜底：LLM 返回的行号可能为 None，确保后续比较/循环不报错
                if data_start_row is None:
                    data_start_row = 1
                if data_end_row is None:
                    data_end_row = ws.max_row

                logger.info(f"  表格数据范围: 第{data_start_row}行 ~ 第{data_end_row}行")
                if table_name:
                    logger.info(f"  检测到表名: '{table_name}' (第{table_name_row}行)")

                # ---- 提取说明文本 ----
                description_before = []
                for row_idx in range(1, data_start_row):
                    if table_name_row and row_idx == table_name_row:
                        continue
                    row_values = []
                    for col_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(row=row_idx, column=col_idx)
                        if cell.value is not None:
                            row_values.append(str(cell.value).strip())
                    if row_values:
                        description_before.append(' '.join(row_values))

                description_after = []
                if data_end_row < ws.max_row:
                    for row_idx in range(data_end_row + 1, ws.max_row + 1):
                        row_values = []
                        for col_idx in range(1, ws.max_column + 1):
                            cell = ws.cell(row=row_idx, column=col_idx)
                            if cell.value is not None:
                                row_values.append(str(cell.value).strip())
                        if row_values:
                            description_after.append(' '.join(row_values))

                metadata[sheet_name] = {
                    'data_start_row': data_start_row,
                    'data_end_row': data_end_row,
                    _TABLE_NAME_KEY: table_name,
                    _TABLE_NAME_ROW_KEY: table_name_row,
                    'description_before': '\n'.join(description_before) if description_before else '',
                    'description_after': '\n'.join(description_after) if description_after else '',
                    'remaining_rows': data_end_row - data_start_row + 1
                }

                # ---- 拆分合并单元格 ----
                fill_count = unmerge_and_fill_sheet(ws)
                logger.info(f"  已拆分合并单元格，填充了 {fill_count} 个空值")

                # ---- 删除说明文本行和表名行 ----
                rows_to_delete = []
                if data_end_row < ws.max_row:
                    for row_idx in range(ws.max_row, data_end_row, -1):
                        rows_to_delete.append(row_idx)
                for row_idx in range(data_start_row - 1, 0, -1):
                    rows_to_delete.append(row_idx)

                if table_name_row and table_name_row not in rows_to_delete:
                    rows_to_delete.append(table_name_row)

                for row_idx in sorted(rows_to_delete, reverse=True):
                    ws.delete_rows(row_idx)

                logger.info(f"  已删除 {len(rows_to_delete)} 行，剩余 {ws.max_row} 行")

                # ---- 检测并合并多行表头 ----
                header_rows = detect_header_rows_with_llm(ws, max_check_rows=8)
                if header_rows and len(header_rows) > 1:
                    logger.info(f"  检测到多行表头: {header_rows}")
                    merged_count = merge_multi_row_header(ws, header_rows)
                elif header_rows and len(header_rows) == 1:
                    logger.info(f"  检测到单行表头: 第{header_rows[0]}行")
                else:
                    logger.info("  未检测到明确的表头行，跳过表头合并")

                # ---- 删除空列 ----
                empty_cols_count = delete_empty_columns(ws)

                # ---- 更新元数据 ----
                metadata[sheet_name]['header_rows_detected'] = header_rows if header_rows else []
                metadata[sheet_name]['header_merged'] = len(header_rows) > 1 if header_rows else False
                metadata[sheet_name]['deleted_empty_cols'] = empty_cols_count

            logger.info("\n处理完成！")

            # ---- 使用内存分块器 ----
            file_name = Path(file_path).stem.replace(" ", "")
            chunker = ExcelChunkerInMemory(wb, metadata, file_name)
            blocks = chunker.chunk_excel()

            # ---- 转换为elements格式 ----
            from langchain_core.documents import Document

            elements = []
            for block in blocks:
                content = block[_CONTENT_KEY]
                meta = block["metadata"]

                chunk_text = content
                doc = Document(
                    page_content=chunk_text,
                    metadata={
                        "source": file_path,
                        "file_name": Path(file_path).stem.replace(" ", ""),
                        "file_type": Path(file_path).suffix.lower()[1:],  # xls 或 xlsx
                        "source_type": "excel",
                        "chunk_index": len(elements) + 1,
                        "total_chunks": len(blocks),
                        "sheet_name": meta.get('sheet_name', ''),
                        "row_number": meta.get('row_number', 0),
                        "row_range": meta.get('row_range', ''),
                        "skipped_sheet": meta.get('skipped_sheet', False),
                    }
                )
                elements.append(doc)

            # ============ 保存 JSON ============
            _save_chunks_to_json(file_path, elements)
            # ==================================

            logger.info(f"\n✅ {file_path} 处理完成，共 {len(elements)} 个元素")
            return file_path, elements

        except Exception as e:
            logger.error(f"❌ {file_path} 处理失败: {e}")
            import traceback
            traceback.print_exc()
            raise

    def _get_elements(self) -> list:
        """
        重写父类方法 - 返回分块后的元素列表
        保持接口不变，外部调用无感知
        """
        if self.max_workers <= 1:
            _, elements = self._process_single_excel(self.file_path)
            return elements

        _, elements = self._process_single_excel(self.file_path)
        return elements


# ============================================================
# 测试代码
# ============================================================

if __name__ == "__main__":
    # 测试串行模式（和原来一样）
    logger.info("=== 串行模式测试 ===")
    _loader = ExcelChineseLoader(
        file_path="D:/任务/海思知识库项目/解析文档/excel/Hi3516CV608 VI输入场景详细说明.xlsx",
        max_workers=1
    )
    docs = _loader.load()
    logger.info(f"加载完成，共 {len(docs)} 个文档片段")
