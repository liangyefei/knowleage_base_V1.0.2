#!/usr/bin/env python3
"""
将 docs/models_api.html（内嵌 OpenAPI 3.1 spec 的 Swagger UI 单文件）
转换为面向客户交付的 Word 接口文档（含完整返回参数说明）。

用法（容器内 / 项目根）：
  python3 tools/html_to_docx.py
  python3 tools/html_to_docx.py -i docs/models_api.html -o docs/models_api.docx

设计要点：
  - 从 HTML 内嵌的 spec JSON 提取结构化数据（不依赖 JS 渲染）
  - 按 tag 分组渲染接口，每个接口含 方法/路径/摘要/请求体/响应 表
  - 每个接口的「响应」部分展开为：
      · 状态码与响应类型
      · 响应结构字段表（含 code/msg/data 各字段说明）
      · 客户端解析提示框（SSE/文件/数组/对象 等不同形态的解析方式）
  - 单独章节渲染所有 Schema（数据模型）
  - 封面 + 信息页 + 目录占位 + 页眉 + 页脚（页码）
  - A4 纸张，中文字体（微软雅黑），代码用 Consolas
"""
# pylint: disable=commented-out-code
import argparse
import json
import logging
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parent.parent

_WFLDCHAR_KEY = 'w:fldChar'
_ALLOF_KEY = 'allOf'

CN_FONT = "Microsoft YaHei"
CN_FONT_BOLD = "Microsoft YaHei"
MONO_FONT = "Consolas"

COLOR_METHOD = {
    "GET": RGBColor(0x19, 0x76, 0xD2),
    "POST": RGBColor(0x38, 0x8E, 0x3C),
    "PUT": RGBColor(0xF5, 0x7C, 0x00),
    "DELETE": RGBColor(0xC6, 0x28, 0x28),
    "PATCH": RGBColor(0x6A, 0x1B, 0x9A),
}
TABLE_HEADER_BG = "D5E8F0"
NOTE_BG = "FFF8E1"
SECTION_BG = "E3F2FD"


# ============================================================
# 通用响应结构字段表（BaseResponse / ListResponse 信封）
# ============================================================

BASE_RESPONSE_FIELDS = [
    ["code", "int", "200", "业务状态码：200 成功；403 拒绝（参数非法）；404 未找到；500 服务异常"],
    ["msg", "string", "success", "状态消息，成功为 success 或具体提示，失败为错误描述"],
    ["data", "object | null", "null", "业务数据，结构随接口变化，详见各接口「响应数据 data 字段」表"],
]

LIST_RESPONSE_FIELDS = [
    ["code", "int", "200", "业务状态码：200 成功；403/404/500 失败"],
    ["msg", "string", "success", "状态消息"],
    ["data", "array<string>", "[]", "字符串数组（知识库名称列表 / 文件名列表）"],
]

DOCUMENT_WITH_VSID_FIELDS = [
    ["id", "string | null", "null", "向量库主键 ID（Milvus pk），由 metadata.id 填充；按 metadata 过滤时可能为 null"],
    ["score", "float", "3.0", "相关性分数，值越小越相关（L2 距离）"],
    ["page_content", "string", "", "文档分块文本内容"],
    ["metadata", "object", "{}", "元数据：source（源文件路径）、content_type、language 等自定义字段"],
    ["type", "string", "Document", "类型标识，固定为 Document"],
]

SOURCE_DOCUMENT_FIELDS = [
    ["text", "string", "", "出处展示文本，含序号、文件名、原文片段"],
    ["num", "int", "1", "出处序号，从 1 开始递增；未找到相关文档时为 0"],
    ["filename", "string", "", "源文件名（如 Hi3798MV310_Datasheet.pdf）"],
    ["url", "string", "", "文档下载链接（仅 multi_knowledge_base_chat 占位项返回，file_chat 不返回此字段）"],
    ["page_content", "string", "", "匹配到的原文分块内容"],
]

FAILED_FILES_FIELDS = [
    ["failed_files", "object", "{}", "失败文件字典，key=文件名，value=错误信息；全部成功时为空对象 {}"],
]

KB_INFO_DATA_FIELDS = [
    ["kb_info", "string", "", "本次设置的知识库介绍文本"],
]

UPLOAD_TEMP_DOCS_DATA_FIELDS = [
    ["id", "string", "", "临时知识库 ID（临时向量库 collection 名），后续 file_chat 接口需传入"],
    ["failed_files", "array<object>", "[]", "上传失败的文件列表，元素为 {文件名: 错误信息}；全部成功时为 []"],
]

SUMMARY_DOC_IDS_DATA_FIELDS = [
    ["summarize", "array<object>", "[]", "生成的摘要文档列表，每个元素为 Document 的 dict 形式（含 page_content/metadata/type）"],
]

MODEL_CONFIG_DATA_FIELDS = [
    ["<模型名>", "object", "", "模型配置字典，含 provider、api_base_url、model_id、device 等字段（已脱敏，不含 api_key）"],
]

LIST_RUNNING_MODELS_DATA_FIELDS = [
    ["<模型名>", "object", "", "在线 LLM 模型配置（已脱敏），key 为模型名称"],
]

LIST_CONFIG_MODELS_DATA_FIELDS = [
    ["local", "object", "", "本地模型配置，key 为模型名，value 为脱敏配置"],
    ["online", "object", "", "在线 API 模型配置，key 为模型名，value 为脱敏配置"],
    ["worker", "object", "", "worker 模型配置（保留字段）"],
]

EMBED_TEXTS_DATA_FIELDS = [
    ["data", "array<array<float>>", "[]", "向量列表，外层对应输入文本顺序，内层为每条文本的 embedding 向量"],
]

SERVER_CONFIGS_FIELDS = [
    ["DEFAULT_KNOWLEDGE_BASE", "string", "", "默认知识库名称"],
    ["DEFAULT_VS_TYPE", "string", "milvus", "默认向量库类型"],
    ["CHUNK_SIZE", "int", "", "文档分块大小"],
    ["OVERLAP_SIZE", "int", "", "分块重叠大小"],
    ["SCORE_THRESHOLD", "float", "", "默认相关性阈值"],
    ["VECTOR_SEARCH_TOP_K", "int", "", "默认向量检索 top_k"],
    ["ZH_TITLE_ENHANCE", "bool", "", "是否开启中文标题加强"],
    ["LLM_MODELS", "array<string>", "[]", "当前配置的 LLM 模型列表"],
    ["HISTORY_LEN", "int", "", "默认历史消息长度"],
    ["TEMPERATURE", "float", "", "默认采样温度"],
    ["EMBEDDING_MODEL", "string", "", "默认 embedding 模型"],
    ["TEXT_SPLITTER_NAME", "string", "", "文本切分器名称"],
    ["api_address", "string", "", "API 服务地址（含 /kw 前缀）"],
]


# ============================================================
# 每个接口的返回参数文档（key 为 (METHOD, PATH) 元组）
#
# 字段说明：
#   - kind: 类型枚举，取值包含 base/list/sse/array/dict/file/str/redirect
#   - data_desc: 对 data 的整体说明（自然语言描述）
#   - data_fields: 字段表格，每行为 [名称, 类型, 默认值, 描述]，无字段时为 None
#   - note: 客户端解析提示（自然语言）
#   - sse_events: SSE 接口专属，每项为 [事件名, 数据结构, 说明]
# ============================================================

RESPONSE_DATA_DOCS = {
    ("GET", "/kw/"): {
        "kind": "redirect",
        "data_desc": "返回 302 重定向至 /kw/docs（Swagger 文档页），无业务数据。",
        "data_fields": None,
        "note": "客户端浏览器直接访问此路径会跳转到 Swagger 文档；如需调用接口，请从 /kw/docs 查看或直接请求具体接口路径。",
    },
    ("POST", "/kw/chat/chat"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出（EventSourceResponse）。每行以 data: 开头，内容为 JSON 字符串。",
        "sse_events": [
            ["流式 token（stream=true）", '{"text": "token片段", "message_id": "xxx"}', "每生成一个 token 即推送一次，message_id 为本次对话记录 ID"],  # pylint: disable=line-too-long
            ["完整回答（stream=false）", '{"text": "完整回答文本", "message_id": "xxx"}', "非流式模式一次性返回完整回答"],
        ],
        "note": "建议使用 stream=true 以提升首字响应体验；message_id 用于后续评分反馈接口 /kw/chat/feedback。SSE 客户端解析示例：onmessage 事件回调中 JSON.parse(e.data) 取 text 字段拼接。",  # pylint: disable=line-too-long
    },
    ("POST", "/kw/chat/next_question"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出，仅推送一条事件。",
        "sse_events": [
            ["预测问题列表", '{"text": ["问题1？", "问题2？", "问题3？"]}', "text 为字符串数组，元素数量等于 predict_number 参数（默认 3）"],
        ],
        "note": "客户端只需接收一条 SSE 事件，JSON.parse(e.data).text 即为预测问题列表。",
    },
    ("POST", "/kw/chat/feedback"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。",
        "data_fields": None,
        "note": "仅通过 code 和 msg 判断反馈是否成功：code=200 表示已记录评分。message_id 来自对话接口（chat / multi_knowledge_base_chat）的 SSE 推送。",
    },
    ("POST", "/kw/chat/file_chat"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出（基于临时上传文件的问答），分两阶段推送。注意：此接口不返回 message_id。",
        "sse_events": [
            ["流式回答 token（stream=true）", '{"answer": "token片段"}', "逐 token 推送回答内容"],
            ["完整回答（stream=false）", '{"answer": "完整回答", "docs": [...]}', "非流式模式一次性返回完整回答与文档出处"],
            ["文档出处", '{"docs": [...]}', "回答结束后推送，docs 为 SourceDocument 数组（不含 url 字段）"],
        ],
        "note": "调用前需先调用 /kw/knowledge_base/upload_temp_docs 上传文件获取 knowledge_id。docs 元素结构见「数据模型 - SourceDocument」（file_chat 的 docs 元素无 url 字段）。",  # pylint: disable=line-too-long
    },
    ("POST", "/kw/llm_model/list_running_models"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为字典：{模型名: 模型配置}。仅返回已配置 api_base_url 的在线 API 模型。",
        "data_fields": LIST_RUNNING_MODELS_DATA_FIELDS,
        "note": "在线 API 模型配置即在线，直接从 ONLINE_LLM_MODEL 返回。返回的配置已脱敏，不含 api_key/secret 等敏感字段。",
    },
    ("POST", "/kw/llm_model/list_config_models"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为按类型分组的模型配置字典。",
        "data_fields": LIST_CONFIG_MODELS_DATA_FIELDS,
        "note": "types 参数控制返回哪些类别（local/online/worker）。每个模型配置均经 get_model_config 脱敏处理。",
    },
    ("POST", "/kw/llm_model/get_model_config"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为指定模型的脱敏配置字典。",
        "data_fields": MODEL_CONFIG_DATA_FIELDS,
        "note": "已自动过滤 worker_class、api_key、secret、以 id 结尾的字段。客户端可直接读取 api_base_url、model_id、provider 等字段。",
    },
    ("POST", "/kw/llm_model/stop"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。",
        "data_fields": None,
        "note": "在线 API 模型配置即在线，无需停止。code=200 + msg 提示即可。",
    },
    ("POST", "/kw/llm_model/change"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。",
        "data_fields": None,
        "note": "在线 API 模型无需切换，需直接修改 params/model_config.yaml 并重启服务。code=200 + msg 提示即可。",
    },
    ("POST", "/kw/server/configs"): {
        "kind": "dict",
        "data_desc": "直接返回配置字典（非 BaseResponse 包装），包含服务运行时配置项。",
        "data_fields": SERVER_CONFIGS_FIELDS,
        "note": "响应体本身就是 JSON 对象，无需解析 code/msg。客户端可直接读取各配置字段。",
    },
    ("POST", "/kw/server/get_prompt_template"): {
        "kind": "str",
        "data_desc": "直接返回 prompt 模板字符串（非 JSON，非 BaseResponse）。",
        "data_fields": None,
        "note": "响应体为纯文本字符串，即为指定 type/name 的 prompt 模板内容。客户端按字符串处理即可。",
    },
    ("POST", "/kw/other/completion"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出，每行 data 为纯文本 token（非 JSON！）。",
        "sse_events": [
            ["流式 token（stream=true）", "纯文本 token 片段", "逐 token 推送，直接拼接即为补全结果"],
            ["完整结果（stream=false）", "纯文本完整结果", "一次性返回完整补全文本"],
        ],
        "note": "与其他 SSE 接口不同，此接口的 data 是纯文本而非 JSON。客户端直接累加 e.data 即可，不要 JSON.parse。",
    },
    ("POST", "/kw/other/embed_texts"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为二维浮点数组（向量列表）。",
        "data_fields": EMBED_TEXTS_DATA_FIELDS,
        "note": "data 外层长度等于输入 texts 长度，内层长度等于 embedding 维度（如 2560）。失败时 code=500，msg 含错误信息。",
    },
    ("GET", "/kw/knowledge_base/list_knowledge_bases"): {
        "kind": "list",
        "data_desc": "ListResponse，data 为知识库名称字符串数组。",
        "data_fields": None,
        "note": "data 元素为知识库名称（如 hi3798mv310、hi3516cv610）。客户端可直接遍历 data 数组。",
    },
    ("POST", "/kw/knowledge_base/create_knowledge_base"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。",
        "data_fields": None,
        "note": "通过 code 和 msg 判断创建结果：200 成功；404 同名已存在或名称为空；403 名称非法；500 创建异常。向量库类型与 embed 模型强制使用默认值。",
    },
    ("POST", "/kw/knowledge_base/delete_knowledge_base"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。删除为异步执行，接口最多等待 5 秒即返回。",
        "data_fields": None,
        "note": "code=200 表示已提交删除任务（实际删除在后台继续）。404 未找到知识库；403 名称非法。",
    },
    ("GET", "/kw/knowledge_base/list_files"): {
        "kind": "list",
        "data_desc": "ListResponse，data 为指定知识库内的文件名字符串数组。",
        "data_fields": None,
        "note": "404 知识库不存在时 data=[]。客户端遍历 data 即可获取文件名列表。",
    },
    ("POST", "/kw/knowledge_base/batch_search_docs"): {
        "kind": "array",
        "data_desc": "直接返回数组（非 BaseResponse），每个元素为 {knowledge_base_name, data: [DocumentWithVSId]}。",
        "data_fields": [
            ["knowledge_base_name", "string", "", "知识库名称"],
            ["data", "array<object>", "[]", "检索到的文档列表，元素结构见「数据模型 - DocumentWithVSId」"],
        ],
        "note": "响应体本身就是 JSON 数组，无需解析 code/msg。每个知识库一个分组，分组内 data 为文档数组（按 language 参数过滤）。",
    },
    ("POST", "/kw/knowledge_base/search_docs"): {
        "kind": "array",
        "data_desc": "直接返回 DocumentWithVSId 数组（非 BaseResponse）。",
        "data_fields": None,
        "note": "响应体为 JSON 数组，元素结构见「数据模型 - DocumentWithVSId」。query 非空时按语义检索；query 为空但 file_name/metadata 非空时按元数据过滤。",
    },
    ("POST", "/kw/knowledge_base/update_docs_by_id"): {
        "kind": "base",
        "data_desc": "BaseResponse，无 data 字段（data=null）。",
        "data_fields": None,
        "note": "通过 code 和 msg 判断更新结果：200 成功；500 知识库不存在或更新失败。",
    },
    ("POST", "/kw/knowledge_base/upload_docs"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "code=200 表示流程完成，但需检查 data.failed_files 是否为空判断是否有文件失败。failed_files 为 {文件名: 错误信息} 字典。",
    },
    ("POST", "/kw/knowledge_base/upload_docs_use_path"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "code=200 表示流程完成，需检查 data.failed_files 判断是否有文件失败。仅支持 SUPPORTED_EXTS 中的文件格式。",
    },
    ("POST", "/kw/knowledge_base/batch_upload_docs"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "按文件夹路径批量入库，数据量大时处理时间较长。code=200 表示流程完成，需检查 data.failed_files 判断是否有文件失败。",
    },
    ("POST", "/kw/knowledge_base/delete_docs"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "code=200 表示删除完成。即使文件不存在也提示成功（前端业务需要）。failed_files 含删除失败的文件及错误信息。",
    },
    ("POST", "/kw/knowledge_base/update_info"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {kb_info} 对象。",
        "data_fields": KB_INFO_DATA_FIELDS,
        "note": "code=200 表示知识库介绍修改完成，data.kb_info 为本次设置的介绍文本。",
    },
    ("POST", "/kw/knowledge_base/update_docs"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "code=200 表示更新完成，data.failed_files 含重新切分/向量化失败的文件。",
    },
    ("POST", "/kw/knowledge_base/update_docs_use_path"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {failed_files} 对象。",
        "data_fields": FAILED_FILES_FIELDS,
        "note": "code=200 表示更新完成，data.failed_files 含失败的文件及错误信息。",
    },
    ("GET", "/kw/knowledge_base/download_doc"): {
        "kind": "file",
        "data_desc": "成功时返回 FileResponse（二进制文件流，Content-Type: multipart/form-data）；失败时返回 BaseResponse。",
        "data_fields": None,
        "note": "preview=true 时 Content-Disposition 为 inline（浏览器内预览）；preview=false 时为 attachment（下载）。客户端按二进制流处理响应体，或直接用浏览器打开。失败时响应为 JSON（BaseResponse），需先判断 Content-Type。",  # pylint: disable=line-too-long
    },
    ("POST", "/kw/knowledge_base/recreate_vector_store"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出处理进度，每个文件处理完推送一次。",
        "sse_events": [
            ["处理进度", '{"code": 200, "msg": "(i/N): filename", "total": N, "finished": i, "doc": "filename"}', "每处理完一个文件推送一次，total 为总文件数，finished 为已完成数"],  # pylint: disable=line-too-long
            ["错误", '{"code": 500, "msg": "添加文件出错：xxx"}', "单个文件处理失败时推送，已跳过继续处理下一个"],
            ["未找到", '{"code": 404, "msg": "未找到知识库 xxx"}', "知识库不存在且 allow_empty_kb=false 时推送"],
        ],
        "note": "客户端可通过 total/finished 显示进度条。SSE 流结束后表示重建完成。",
    },
    ("POST", "/kw/knowledge_base/upload_temp_docs"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {id, failed_files} 对象。",
        "data_fields": UPLOAD_TEMP_DOCS_DATA_FIELDS,
        "note": "data.id 为临时知识库 ID，后续 /kw/chat/file_chat 接口需传入此 ID。failed_files 为 [{文件名: 错误信息}] 列表。",
    },
    ("POST", "/kw/knowledge_base/kb_summary_api/summary_file_to_vector_store"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出，按文件名生成摘要后推送一次。",
        "sse_events": [
            ["摘要完成", '{"code": 200, "msg": "filename 总结完成", "doc": "filename"}', "单个文件摘要生成完成时推送"],
            ["错误", '{"code": 500, "msg": "知识库总结文件 xxx 时出错。已跳过。"}', "摘要生成失败时推送"],
            ["未找到", '{"code": 404, "msg": "未找到知识库 xxx"}', "知识库不存在且 allow_empty_kb=false 时推送"],
        ],
        "note": "客户端通过 code 判断每个文件的处理结果。SSE 流结束后表示摘要生成完成。",
    },
    ("POST", "/kw/knowledge_base/kb_summary_api/summary_doc_ids_to_vector_store"): {
        "kind": "base",
        "data_desc": "BaseResponse，data 为 {summarize} 对象。",
        "data_fields": SUMMARY_DOC_IDS_DATA_FIELDS,
        "note": "code=200 表示总结完成，data.summarize 为生成的摘要文档列表。404 知识库不存在时 data={}。",
    },
    ("POST", "/kw/knowledge_base/kb_summary_api/recreate_summary_vector_store"): {
        "kind": "sse",
        "data_desc": "SSE 流式输出处理进度，每个文件摘要生成完推送一次。",
        "sse_events": [
            ["处理进度", '{"code": 200, "msg": "(i/N): filename", "total": N, "finished": i, "doc": "filename"}', "每生成完一个文件摘要推送一次"],  # pylint: disable=line-too-long
            ["错误", '{"code": 500, "msg": "知识库总结文件 xxx 时出错。已跳过。"}', "单个文件摘要生成失败时推送"],
            ["未找到", '{"code": 404, "msg": "未找到知识库 xxx"}', "知识库不存在且 allow_empty_kb=false 时推送"],
        ],
        "note": "客户端可通过 total/finished 显示进度条。SSE 流结束后表示重建完成。此接口会先清空旧摘要再重建。",
    },
}


# ============================================================
# OpenAPI spec 提取与解析
# ============================================================

def extract_spec(html_path: Path) -> dict:
    html = html_path.read_text(encoding="utf-8")
    m = re.search(r'spec:\s*(\{"openapi".*?\}),\s*dom_id', html, re.S)
    if not m:
        raise RuntimeError("未在 HTML 中找到 spec JSON")
    spec_str = m.group(1).replace("<\\/", "/")
    return json.loads(spec_str)


def resolve_ref(spec: dict, ref: str):
    if not ref:
        return None
    parts = ref.lstrip("#/").split("/")
    node = spec
    for p in parts:
        if not isinstance(node, dict):
            return None
        node = node.get(p, {})
    return node or None


def get_resolved_schema(spec: dict, schema):
    if not schema:
        return None
    if "$ref" in schema:
        return resolve_ref(spec, schema["$ref"])
    if _ALLOF_KEY in schema:
        merged = {}
        for sub in schema[_ALLOF_KEY]:
            sub_r = get_resolved_schema(spec, sub) or {}
            merged.update(sub_r)
        for k, v in schema.items():
            if k != _ALLOF_KEY:
                merged[k] = v
        return merged
    return schema


def schema_name_from_ref(ref: str) -> str:
    if not ref:
        return ""
    return ref.rstrip("/").split("/")[-1]


def type_str(schema) -> str:
    if not schema:
        return ""
    if "$ref" in schema:
        return schema_name_from_ref(schema["$ref"])
    if "type" in schema:
        t = schema["type"]
        if t == "array":
            inner = type_str(schema.get("items", {}))
            return f"array<{inner}>"
        if isinstance(t, list):
            return " | ".join(t)
        return str(t)
    if "anyOf" in schema:
        return " | ".join(type_str(s) for s in schema["anyOf"])
    if "enum" in schema:
        return "enum"
    return ""


# ============================================================
# docx 渲染辅助
# ============================================================

def set_cell_shading(cell, color_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()  # pylint: disable=protected-access
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def set_cell_text(cell, text, **kwargs):
    """设置单元格文本及样式。

    可选参数（通过 kwargs 传入，避免签名参数过多）:
        bold: bool - 是否加粗（默认 False）
        font: str - 字体（默认 CN_FONT）
        size: int/float - 字号（默认 10）
        color: RGBColor or None - 字体颜色（默认 None）
        align_left: bool - 是否左对齐（默认 True，False 则居中）
    """
    bold = kwargs.get("bold", False)
    font = kwargs.get("font", CN_FONT)
    size = kwargs.get("size", 10)
    color = kwargs.get("color")
    align_left = kwargs.get("align_left", True)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if align_left else WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text) if text is not None else "")
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()  # pylint: disable=protected-access
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), CN_FONT)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    if color is not None:
        run.font.color.rgb = color


def set_table_borders(table):
    tbl = table._tbl  # pylint: disable=protected-access
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "B0B0B0")
        borders.append(b)
    tblPr.append(borders)


def add_table(doc, headers, rows, col_widths_cm=None, mono_cols=None):
    n_cols = len(headers)
    mono_cols = mono_cols or set()
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    if col_widths_cm:
        for i, w in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, TABLE_HEADER_BG)
        set_cell_text(cell, h, bold=True, size=10)
    for r_idx, row_data in enumerate(rows, start=1):
        for c_idx, val in enumerate(row_data):
            cell = table.rows[r_idx].cells[c_idx]
            mono = c_idx in mono_cols or (c_idx == 0 and val and any(ch in str(val) for ch in "/<>(){}."))
            set_cell_text(
                cell,
                val if val is not None else "",
                font=MONO_FONT if mono else CN_FONT,
                size=9 if mono else 10,
            )
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    return table


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)  # pylint: disable=protected-access
    run._r.append(instr)  # pylint: disable=protected-access
    run._r.append(fld_end)  # pylint: disable=protected-access


def style_run(run, font=CN_FONT, size=10.5, bold=False, color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()  # pylint: disable=protected-access
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), CN_FONT)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    run = p.add_run(text)
    sizes = {1: 18, 2: 14, 3: 12}
    style_run(run, font=CN_FONT_BOLD, size=sizes.get(level, 11), bold=True)
    return p


def add_paragraph_text(doc, text, **kwargs):
    """添加段落文本，可指定样式。

    可选参数（通过 kwargs 传入，避免签名参数过多）:
        size: float - 字号（默认 10.5）
        bold: bool - 是否加粗（默认 False）
        mono: bool - 是否使用等宽字体（默认 False）
        color: RGBColor or None - 字体颜色（默认 None）
    """
    size = kwargs.get("size", 10.5)
    bold = kwargs.get("bold", False)
    mono = kwargs.get("mono", False)
    color = kwargs.get("color")
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_run(run, font=MONO_FONT if mono else CN_FONT, size=size, bold=bold, color=color)
    return p


def add_note_box(doc, text, title="客户端解析提示"):
    """添加一个带背景色的提示框（单行表格实现）。"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(16.0)
    set_cell_shading(cell, NOTE_BG)
    cell.text = ""
    p = cell.paragraphs[0]
    run_t = p.add_run(f"【{title}】")
    style_run(run_t, font=CN_FONT_BOLD, size=10, bold=True, color=RGBColor(0xE6, 0x6C, 0x00))
    run_c = p.add_run("  " + text)
    style_run(run_c, font=CN_FONT, size=10)
    doc.add_paragraph()
    return table


# ============================================================
# 请求体渲染
# ============================================================

def render_request_body(doc, spec, body: dict):
    if not body:
        return
    content = body.get("content", {})
    app_json = content.get("application/json") or content.get("application/x-www-form-urlencoded") or content.get("multipart/form-data")  # pylint: disable=line-too-long
    if not app_json:
        return
    schema_ref = app_json.get("schema", {})
    schema = get_resolved_schema(spec, schema_ref)
    if not schema:
        return
    required = set(schema.get("required", []) or [])
    props = schema.get("properties", {}) or {}
    if not props:
        add_paragraph_text(doc, f"请求体 schema: {schema_name_from_ref(schema_ref.get('$ref', ''))}（无 properties）", size=9, color=RGBColor(0x60, 0x60, 0x60))  # pylint: disable=line-too-long
        return
    rows = []
    for name, sub in props.items():
        rows.append([
            name,
            type_str(sub),
            "是" if name in required else "否",
            sub.get("default", "") if sub.get("default", "") is not None else "",
            sub.get("description", "") or "",
        ])
    add_table(doc, ["字段", "类型", "必填", "默认值", "说明"], rows,
              col_widths_cm=[3.2, 2.8, 1.2, 2.5, 6.3], mono_cols={0})


# ============================================================
# 响应渲染（核心：展开返回参数详细说明）
# ============================================================

def render_response_kind_badge(doc, kind: str):
    """响应类型标签。"""
    kind_label = {
        "base": "BaseResponse（标准信封）",
        "list": "ListResponse（列表信封）",
        "sse": "SSE 流式输出（EventSourceResponse）",
        "array": "直接返回数组",
        "dict": "直接返回对象",
        "file": "文件流（FileResponse）",
        "str": "纯文本字符串",
        "redirect": "HTTP 重定向",
    }.get(kind, kind)
    p = doc.add_paragraph()
    run1 = p.add_run("响应类型：")
    style_run(run1, font=CN_FONT_BOLD, size=10, bold=True)
    run2 = p.add_run(kind_label)
    style_run(run2, font=CN_FONT, size=10, color=RGBColor(0x15, 0x6C, 0xC2))


def render_response_for_endpoint(doc, spec, path: str, method: str, responses: dict):
    """渲染单个接口的响应部分：状态码表 + 信封字段表 + data 字段表 + SSE 事件表 + 解析提示。"""
    _RENDER_RESPONSE_LABEL = '说明'
    _BASE_KIND = 'base'
    _FIELD_LABEL = '字段'
    _TYPE_LABEL = '类型'
    _RESPONSE_BODY_LABEL = '响应体'
    if not responses:
        return

    doc_key = (method.upper(), path)
    rdoc = RESPONSE_DATA_DOCS.get(doc_key)
    if rdoc is None:
        add_paragraph_text(doc, "（未收录此接口的返回参数文档，仅展示 OpenAPI 默认响应表）", size=9, color=RGBColor(0x60, 0x60, 0x60))
        rows = []
        for code, resp in responses.items():
            rows.append([code, resp.get("description", ""), ""])
        add_table(doc, ["状态码", _RENDER_RESPONSE_LABEL, "返回 schema"], rows, col_widths_cm=[2.0, 8.5, 5.5])
        return

    kind = rdoc.get("kind", _BASE_KIND)

    render_response_kind_badge(doc, kind)

    add_paragraph_text(doc, "响应说明", size=10, bold=True)
    add_paragraph_text(doc, rdoc.get("data_desc", ""), size=10)

    if kind in (_BASE_KIND, "list"):
        env_rows = BASE_RESPONSE_FIELDS if kind == _BASE_KIND else LIST_RESPONSE_FIELDS
        add_paragraph_text(doc, "响应结构（信封字段）", size=10, bold=True)
        add_table(doc, [_FIELD_LABEL, _TYPE_LABEL, "默认值", _RENDER_RESPONSE_LABEL], env_rows,
                  col_widths_cm=[2.8, 3.0, 2.2, 8.0], mono_cols={0})

        data_fields = rdoc.get("data_fields")
        if data_fields:
            add_paragraph_text(doc, "响应数据 data 字段", size=10, bold=True)
            add_table(doc, [_FIELD_LABEL, _TYPE_LABEL, "默认值/示例", _RENDER_RESPONSE_LABEL], data_fields,
                      col_widths_cm=[3.5, 3.0, 2.5, 7.0], mono_cols={0})

    elif kind == "sse":
        sse_events = rdoc.get("sse_events") or []
        if sse_events:
            add_paragraph_text(doc, "SSE 事件结构", size=10, bold=True)
            add_table(doc, ["事件", "data 结构", _RENDER_RESPONSE_LABEL], sse_events,
                      col_widths_cm=[3.5, 6.0, 6.5], mono_cols={1})

    elif kind in ("array", "dict"):
        data_fields = rdoc.get("data_fields")
        if data_fields:
            add_paragraph_text(doc, "响应体字段（直接返回，非信封包装）", size=10, bold=True)
            add_table(doc, [_FIELD_LABEL, _TYPE_LABEL, "默认值/示例", _RENDER_RESPONSE_LABEL], data_fields,
                      col_widths_cm=[3.5, 3.0, 2.5, 7.0], mono_cols={0})

    elif kind == "file":
        add_paragraph_text(doc, _RESPONSE_BODY_LABEL, size=10, bold=True)
        add_paragraph_text(doc, "成功：HTTP 200，Content-Type: multipart/form-data，响应体为文件二进制流。", size=10)
        add_paragraph_text(doc, "失败：HTTP 200，Content-Type: application/json，响应体为 BaseResponse（{code, msg, data}）。", size=10)  # pylint: disable=line-too-long

    elif kind == "str":
        add_paragraph_text(doc, _RESPONSE_BODY_LABEL, size=10, bold=True)
        add_paragraph_text(doc, "HTTP 200，Content-Type: text/plain; charset=utf-8，响应体为纯文本字符串。", size=10)

    elif kind == "redirect":
        add_paragraph_text(doc, _RESPONSE_BODY_LABEL, size=10, bold=True)
        add_paragraph_text(doc, "HTTP 302，Location: /kw/docs，无业务响应体。", size=10)

    note = rdoc.get("note")
    if note:
        add_note_box(doc, note)


# ============================================================
# 接口渲染
# ============================================================

def render_endpoint(doc, spec, path: str, method: str, op: dict):
    p = doc.add_paragraph()
    method_upper = method.upper()
    color = COLOR_METHOD.get(method_upper, RGBColor(0x33, 0x33, 0x33))
    run_m = p.add_run(f"{method_upper}  ")
    style_run(run_m, font=MONO_FONT, size=11, bold=True, color=color)
    run_p = p.add_run(path)
    style_run(run_p, font=MONO_FONT, size=11, bold=True)
    summary = op.get("summary", "")
    if summary:
        p2 = doc.add_paragraph()
        run_s = p2.add_run(summary)
        style_run(run_s, font=CN_FONT, size=11, bold=True)
    desc = op.get("description", "")
    if desc:
        add_paragraph_text(doc, desc, size=10)
    op_id = op.get("operationId", "")
    if op_id:
        add_paragraph_text(doc, f"Operation ID: {op_id}", size=9, mono=True, color=RGBColor(0x60, 0x60, 0x60))
    params = op.get("parameters", [])
    if params:
        add_paragraph_text(doc, "路径/查询参数", size=10, bold=True)
        rows = []
        for prm in params:
            sch = prm.get("schema", {})
            rows.append([
                prm.get("name", ""),
                prm.get("in", ""),
                type_str(sch),
                "是" if prm.get("required") else "否",
                prm.get("description", "") or "",
            ])
        add_table(doc, ["参数", "位置", "类型", "必填", "说明"], rows,
                  col_widths_cm=[3.2, 1.8, 2.5, 1.2, 7.3], mono_cols={0})
    body = op.get("requestBody")
    if body:
        add_paragraph_text(doc, "请求体", size=10, bold=True)
        render_request_body(doc, spec, body)
    responses = op.get("responses", {})
    if responses:
        add_paragraph_text(doc, "响应", size=10, bold=True)
        render_response_for_endpoint(doc, spec, path, method, responses)


def render_schema(doc, spec, name: str, schema: dict):
    add_heading(doc, name, level=3)
    desc = schema.get("description", "")
    if desc:
        add_paragraph_text(doc, desc, size=10)
    required = set(schema.get("required", []) or [])
    props = schema.get("properties", {}) or {}
    if not props:
        add_paragraph_text(doc, "（无 properties）", size=9, color=RGBColor(0x60, 0x60, 0x60))
        return
    rows = []
    for pname, psub in props.items():
        rows.append([
            pname,
            type_str(psub),
            "是" if pname in required else "否",
            psub.get("default", "") if psub.get("default", "") is not None else "",
            psub.get("description", "") or "",
        ])
    add_table(doc, ["字段", "类型", "必填", "默认值", "说明"], rows,
              col_widths_cm=[3.2, 2.8, 1.2, 2.5, 6.3], mono_cols={0})


# ============================================================
# 文档章节构建
# ============================================================

def build_cover(doc, info: dict, n_paths: int, n_schemas: int):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(info.get("title", "API Server"))
    style_run(run, font=CN_FONT_BOLD, size=32, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("API 接口文档")
    style_run(run2, font=CN_FONT, size=18)
    for _ in range(2):
        doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(f"版本 {info.get('version', '')}")
    style_run(run3, font=CN_FONT, size=14)
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run4 = p4.add_run(f"发布日期 {date.today().strftime('%Y 年 %m 月 %d 日')}")
    style_run(run4, font=CN_FONT, size=12)
    for _ in range(2):
        doc.add_paragraph()
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run5 = p5.add_run(f"共 {n_paths} 个接口  ·  {n_schemas} 个数据模型")
    style_run(run5, font=CN_FONT, size=12, color=RGBColor(0x60, 0x60, 0x60))
    for _ in range(6):
        doc.add_paragraph()
    p6 = doc.add_paragraph()
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run6 = p6.add_run("本文档仅供客户内部评估使用，未经许可不得外传")
    style_run(run6, font=CN_FONT, size=10, color=RGBColor(0x90, 0x90, 0x90))
    doc.add_page_break()


def build_info_section(doc, spec: dict):
    info = spec.get("info", {})
    add_heading(doc, "1. 文档概述", level=1)
    desc = info.get("description", "")
    if desc:
        for line in desc.split("\n"):
            line = line.strip()
            if line:
                add_paragraph_text(doc, line, size=10.5)
    add_paragraph_text(doc, f"接口版本：{info.get('version', '')}", size=10.5)
    add_paragraph_text(doc, "Base URL：http://<服务器地址>:7861", size=10.5, mono=True)
    add_paragraph_text(doc, "文档入口：GET /kw/（跳转至 Swagger 文档）", size=10.5)
    doc.add_page_break()


def build_toc_page(doc):
    add_heading(doc, "2. 目录", level=1)
    p = doc.add_paragraph()
    run = p.add_run("（在 Word 中按 Ctrl+A 全选后按 F9 可更新此目录页字段）")
    style_run(run, font=CN_FONT, size=9, color=RGBColor(0x90, 0x90, 0x90))
    para = doc.add_paragraph()
    _WFLDCHARTYPE_ATTR = 'w:fldCharType'
    fld_begin = OxmlElement(_WFLDCHAR_KEY)
    fld_begin.set(qn(_WFLDCHARTYPE_ATTR), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement(_WFLDCHAR_KEY)
    fld_sep.set(qn(_WFLDCHARTYPE_ATTR), "separate")
    fld_end = OxmlElement(_WFLDCHAR_KEY)
    fld_end.set(qn(_WFLDCHARTYPE_ATTR), "end")
    run2 = para.add_run()
    run2._r.append(fld_begin)  # pylint: disable=protected-access
    run2._r.append(instr)  # pylint: disable=protected-access
    run2._r.append(fld_sep)  # pylint: disable=protected-access
    run2._r.append(fld_end)  # pylint: disable=protected-access
    doc.add_page_break()


def build_envelope_overview(doc):
    """第 3 章：响应约定与通用响应结构。"""
    add_heading(doc, "3. 响应约定", level=1)
    add_paragraph_text(doc, "本服务接口的响应形态分为以下几类，客户端需按不同方式解析：", size=10.5)

    add_heading(doc, "3.1 BaseResponse（标准信封）", level=2)
    add_paragraph_text(doc, "大多数接口返回 BaseResponse，结构如下：", size=10.5)
    add_table(doc, ["字段", "类型", "默认值", "说明"], BASE_RESPONSE_FIELDS,
              col_widths_cm=[2.8, 3.0, 2.2, 8.0], mono_cols={0})
    add_paragraph_text(doc, "客户端解析流程：先判断 code == 200，再按接口文档解析 data 字段。", size=10)
    add_note_box(doc, "data 字段的结构因接口而异（可能是对象、数组、null），请务必查阅对应接口的「响应数据 data 字段」表。", title="data 字段形态")

    add_heading(doc, "3.2 ListResponse（列表信封）", level=2)
    add_paragraph_text(doc, "列表类接口（如获取知识库列表、文件列表）返回 ListResponse，data 固定为字符串数组：", size=10.5)
    add_table(doc, ["字段", "类型", "默认值", "说明"], LIST_RESPONSE_FIELDS,
              col_widths_cm=[2.8, 3.0, 2.2, 8.0], mono_cols={0})

    add_heading(doc, "3.3 SSE 流式输出（EventSourceResponse）", level=2)
    add_paragraph_text(doc, "对话类与重建类接口采用 SSE（Server-Sent Events）流式输出，Content-Type: text/event-stream。", size=10.5)
    add_paragraph_text(doc, "每行格式为 data: <payload>\\n\\n，payload 一般为 JSON 字符串（/kw/other/completion 除外，其 payload 为纯文本）。", size=10.5)  # pylint: disable=line-too-long
    add_paragraph_text(doc, "客户端解析流程：", size=10, bold=True)
    add_paragraph_text(doc, "1. 使用 EventSource（浏览器）或 httpx/requests 流式读取（后端）建立连接；", size=10)
    add_paragraph_text(doc, "2. 按行读取 data: 前缀后的内容；", size=10)
    add_paragraph_text(doc, "3. JSON.parse 后取对应字段（如 text/answer/docs/message_id）；", size=10)
    add_paragraph_text(doc, "4. 流结束（连接关闭）表示本次响应完成。", size=10)
    add_note_box(doc, "/kw/other/completion 的 SSE data 为纯文本 token，不要 JSON.parse，直接拼接即可。", title="例外接口")

    add_heading(doc, "3.4 直接返回数组/对象", level=2)
    add_paragraph_text(doc, "部分检索接口（如 search_docs、batch_search_docs）直接返回 JSON 数组或对象，无 BaseResponse 信封包装。", size=10.5)
    add_paragraph_text(doc, "客户端直接按数组/对象解析响应体即可，无需判断 code/msg。", size=10)

    add_heading(doc, "3.5 文件流（FileResponse）", level=2)
    add_paragraph_text(doc, "download_doc 接口成功时返回二进制文件流（Content-Type: multipart/form-data），失败时返回 BaseResponse（JSON）。", size=10.5)  # pylint: disable=line-too-long
    add_paragraph_text(doc, "客户端应先判断响应 Content-Type：非 JSON 时按二进制流保存文件，JSON 时解析为 BaseResponse 处理错误。", size=10)

    add_heading(doc, "3.6 纯文本与重定向", level=2)
    add_paragraph_text(doc, "get_prompt_template 接口直接返回纯文本字符串（Content-Type: text/plain）；GET /kw/ 返回 302 重定向至 /kw/docs。", size=10.5)  # pylint: disable=line-too-long
    doc.add_page_break()


def build_endpoint_index(doc, spec: dict):
    add_heading(doc, "4. 接口总览", level=1)
    rows = []
    for path, ops in spec.get("paths", {}).items():
        for method, op in ops.items():
            tag = (op.get("tags") or ["其他"])[0]
            rows.append([method.upper(), path, op.get("summary", ""), tag])
    add_table(doc, ["方法", "路径", "说明", "分组"], rows,
              col_widths_cm=[1.8, 5.5, 6.2, 2.5], mono_cols={1})
    doc.add_page_break()


def build_endpoint_details(doc, spec: dict):
    add_heading(doc, "5. 接口详情", level=1)
    add_paragraph_text(doc, "以下按 tag 分组列出所有接口的请求/响应定义。每个接口的「响应」部分含响应类型、信封字段、data 字段、SSE 事件结构与客户端解析提示。", size=10.5)
    groups = {}
    for path, ops in spec.get("paths", {}).items():
        for method, op in ops.items():
            tag = (op.get("tags") or ["其他"])[0]
            groups.setdefault(tag, []).append((path, method, op))
    sec_idx = 0
    for tag, items in groups.items():
        sec_idx += 1
        add_heading(doc, f"5.{sec_idx} {tag}", level=2)
        for path, method, op in items:
            render_endpoint(doc, spec, path, method, op)
            doc.add_paragraph()
        doc.add_page_break()


def build_schemas_section(doc, spec: dict):
    add_heading(doc, "6. 数据模型", level=1)
    add_paragraph_text(doc, "以下列出关键数据模型。OpenAPI 自动生成的 Pydantic schema 已含请求体与响应模型，此处补充几个客户端解析时常用的复合结构。", size=10.5)

    add_heading(doc, "6.1 DocumentWithVSId（检索结果文档）", level=2)
    add_paragraph_text(doc, "search_docs / batch_search_docs 返回的文档结构，继承自 langchain Document，增加 id 与 score 字段。", size=10)
    add_table(doc, ["字段", "类型", "默认值", "说明"], DOCUMENT_WITH_VSID_FIELDS,
              col_widths_cm=[3.2, 3.0, 2.5, 7.3], mono_cols={0})

    add_heading(doc, "6.2 SourceDocument（对话出处文档）", level=2)
    add_paragraph_text(doc, "multi_knowledge_base_chat / file_chat 接口 SSE 推送的 docs 数组元素结构。", size=10)
    add_table(doc, ["字段", "类型", "默认值", "说明"], SOURCE_DOCUMENT_FIELDS,
              col_widths_cm=[3.2, 3.0, 2.5, 7.3], mono_cols={0})
    add_paragraph_text(doc, "注：file_chat 接口的 SourceDocument 不含 url 字段；multi_knowledge_base_chat 占位项含 url（文档下载链接，正常 docs 无此字段）。", size=9, color=RGBColor(0x60, 0x60, 0x60))  # pylint: disable=line-too-long

    add_heading(doc, "6.3 OpenAPI 自动生成的 Schema", level=2)
    schemas = spec.get("components", {}).get("schemas", {}) or {}
    for name, sch in sorted(schemas.items()):
        render_schema(doc, spec, name, sch)
    doc.add_page_break()


def build_change_log(doc):
    add_heading(doc, "7. 变更记录", level=1)
    rows = [
        [date.today().strftime("%Y-%m-%d"), "v0.2.10", "首次生成客户交付版接口文档"],
        [date.today().strftime("%Y-%m-%d"), "v0.2.10", "补充各接口返回参数详细说明、SSE 事件结构、客户端解析提示"],
    ]
    add_table(doc, ["日期", "版本", "说明"], rows, col_widths_cm=[3.0, 2.5, 10.5])


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_p.add_run("HaiSI RAG API Server  ·  接口文档")
    style_run(run, font=CN_FONT, size=9, color=RGBColor(0x90, 0x90, 0x90))
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = footer_p.add_run("第 ")
    style_run(run1, font=CN_FONT, size=9, color=RGBColor(0x90, 0x90, 0x90))
    add_page_number_field(footer_p)
    run2 = footer_p.add_run(" 页")
    style_run(run2, font=CN_FONT, size=9, color=RGBColor(0x90, 0x90, 0x90))
    for s in doc.styles:
        if s.name and s.name.startswith("Heading"):
            try:
                s.font.name = CN_FONT_BOLD
                s.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_BOLD)
            except Exception as e:
                logger.debug("[html-to-docx] failed to set Heading font for style %s: %s", s.name, e)
                pass


def main():
    ap = argparse.ArgumentParser(description="将 models_api.html 转为客户交付用 Word 文档（含返回参数说明）")
    ap.add_argument("-i", "--input", default="docs/models_api.html", help="输入 HTML 文件")
    ap.add_argument("-o", "--output", default="docs/models_api.docx", help="输出 docx 文件")
    args = ap.parse_args()

    html_path = (ROOT / args.input) if not Path(args.input).is_absolute() else Path(args.input)
    out_path = (ROOT / args.output) if not Path(args.output).is_absolute() else Path(args.output)

    logging.info("[1/4] 提取 spec: %s", html_path)
    spec = extract_spec(html_path)
    info = spec.get("info", {})
    n_paths = sum(len(o) for o in spec.get("paths", {}).values())
    n_schemas = len(spec.get("components", {}).get("schemas", {}) or {})
    logging.info("     标题=%s, 版本=%s, 接口=%s, 模型=%s", info.get('title'), info.get('version'), n_paths, n_schemas)

    logging.info("[2/4] 构建 Word 文档结构")
    doc = Document()
    setup_page(doc)
    build_cover(doc, info, n_paths, n_schemas)
    build_info_section(doc, spec)
    build_toc_page(doc)
    build_envelope_overview(doc)
    build_endpoint_index(doc, spec)
    build_endpoint_details(doc, spec)
    build_schemas_section(doc, spec)
    build_change_log(doc)

    logging.info("[3/4] 写出: %s", out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))

    logging.info("[4/4] 完成。文件大小: %s KB", out_path.stat().st_size // 1024)


if __name__ == "__main__":
    main()
